"""BatchWatch web app: check a pack or a stock list against the alert registry, browse alerts, watch the watcher.

Run locally:  uvicorn app.main:app --port 8000
Uploaded stock files are parsed in memory and never written to disk.
"""
import io
import json
import os
import queue
import threading
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv

from fastapi import FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles

from app.limits import DailyLimiter
from app.store import store
from app.views import alert_summary, decorate_trace, insights, match_json, registry_stats, search, status
from batchwatch.gs1 import parse_gs1
from batchwatch.match import REQUIRED_COLUMNS, lookup_pack, map_columns, match_inventory, parse_inventory, read_rows
from batchwatch.notify import atom_feed
from batchwatch.pipeline import SEED
from batchwatch.trace import trace_cached

ROOT = Path(__file__).resolve().parent.parent
load_dotenv(ROOT / ".env")
_trace_lock = threading.Lock()  # one agent run at a time keeps API spend predictable
limiter = DailyLimiter.from_env()
SAMPLES = {"riyadh": ROOT / "examples" / "riyadh_pharmacy.csv", "hospital": ROOT / "examples" / "inventory_sample.csv"}
MAX_UPLOAD_BYTES = 2_000_000
MAX_LINES = 20_000

app = FastAPI(title="BatchWatch", docs_url="/api/docs", redoc_url=None)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def entries() -> list[dict]:
    return store.entries()


def _by_id() -> dict:
    return {e["alert_id"]: e for e in entries()}


def alerts_as_of() -> str | None:
    """When the registry was last brought up to date: the last watcher run, else the newest extraction."""
    runs = store.runs(1)
    if runs and runs[-1].get("finished"):
        return runs[-1]["finished"]
    return max((e.get("extracted_at") or "" for e in entries()), default=None) or None


def _matches(ms) -> list[dict]:
    by_id = _by_id()
    return [match_json(m, by_id.get(m.alert.alert_id), store.trace(m.alert.alert_id)) for m in ms]


# ---------------------------------------------------------------- stock lists

def _decode(raw: bytes) -> str:
    """UTF-8 first; for legacy exports, Arabic Windows (cp1256) when the text looks Arabic, else Western (cp1252)."""
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    arabic = raw.decode("cp1256", errors="replace")
    if sum(1 for ch in arabic[:4000] if "؀" <= ch <= "ۿ") >= 3:
        return arabic
    try:
        return raw.decode("cp1252")
    except UnicodeDecodeError:
        raise HTTPException(400, "Could not read the file. Save it as CSV (UTF-8) or Excel and try again.")


def _xlsx_to_csv(raw: bytes) -> str:
    import csv

    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "Could not open the Excel file. Save it as .xlsx or CSV and try again.")
    out = io.StringIO()
    writer = csv.writer(out)
    for row in wb.worksheets[0].iter_rows(values_only=True):
        if any(v not in (None, "") for v in row):
            writer.writerow(["" if v is None else (v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else v) for v in row])
    return out.getvalue()


def _check(text: str, filename: str) -> dict:
    headers, rows = read_rows(text)
    if not headers:
        raise HTTPException(400, "The file is empty.")
    mapping, how = map_columns(headers), "names"
    if any(f not in mapping for f in REQUIRED_COLUMNS) and os.environ.get("NEBIUS_API_KEY"):
        try:  # an unfamiliar export: let Nemotron read the header row and three sample rows
            from batchwatch.columns import suggest_columns
            from batchwatch.llm import make_client
            suggested = suggest_columns(make_client(), headers, rows)
            if all(f in suggested for f in REQUIRED_COLUMNS):
                mapping, how = suggested, "nemotron"
        except Exception:
            pass
    missing = [f for f in REQUIRED_COLUMNS if f not in mapping]
    if missing:
        raise HTTPException(400, "Couldn't find a product column and a batch column. Your columns are: "
                                 + ", ".join(headers[:20]) + ". Rename them (for example Product, Batch) or use the template.")
    items = parse_inventory(text, mapping)
    if not items:
        raise HTTPException(400, "The file has a header but no stock lines.")
    if len(items) > MAX_LINES:
        raise HTTPException(400, f"The file has {len(items)} lines; the limit is {MAX_LINES}.")
    matches = match_inventory(items, entries())
    counts = Counter(m.tier for m in matches)
    return {
        "file": filename,
        "lines_checked": len(items),
        "alerts_checked": len(entries()),
        "checked_at": _now(),
        "alerts_as_of": alerts_as_of(),
        "columns": mapping,
        "columns_by": how,
        "counts": {t: counts.get(t, 0) for t in ("CONFIRMED", "VERIFY", "WATCH", "WITHDRAWN")},
        "matches": _matches(matches),
    }


@app.post("/api/check")
async def post_check(file: UploadFile = File(...)):
    raw = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(400, "The file is larger than 2 MB. Split it or export only the medicines in stock.")
    name = file.filename or "stock.csv"
    text = _xlsx_to_csv(raw) if raw[:2] == b"PK" or name.lower().endswith(".xlsx") else _decode(raw)
    return _check(text, name)


@app.get("/api/check/sample")
def get_sample_check(name: str = Query("riyadh", pattern="^(riyadh|hospital)$")):
    path = SAMPLES[name]
    return _check(path.read_text(encoding="utf-8-sig"), path.name)


@app.get("/api/sample.csv")
def get_sample_csv():
    return FileResponse(SAMPLES["hospital"], media_type="text/csv", filename="batchwatch_stock_template.csv")


# ---------------------------------------------------------------- one pack

@app.get("/api/lookup")
def get_lookup(code: str = Query("", max_length=200), batch: str = Query("", max_length=40),
               gtin: str = Query("", max_length=20), expiry: str = Query("", max_length=20),
               product: str = Query("", max_length=120)):
    """Check one pack: a scanned GS1 DataMatrix string and/or typed fields (typed fields win). No AI, no cost."""
    parsed = parse_gs1(code) if code else {}
    fields = {"gtin": gtin or parsed.get("gtin", ""), "batch": batch or parsed.get("batch", ""),
              "expiry": expiry or parsed.get("expiry", ""), "serial": parsed.get("serial", ""), "product": product}
    if not fields["batch"]:
        raise HTTPException(400, "No batch number found. Scan the 2D code on the pack or type the batch number.")
    matches = lookup_pack(entries(), fields["batch"], fields["gtin"], product, fields["expiry"])
    return {"pack": fields, "alerts_checked": len(entries()), "checked_at": _now(), "alerts_as_of": alerts_as_of(),
            "matches": _matches(matches)}


# ---------------------------------------------------------------- alerts

@app.get("/api/registry")
def get_registry():
    rows = sorted(entries(), key=lambda e: str(e.get("date", "")).replace("-", ""), reverse=True)
    return {"stats": registry_stats(rows), "alerts": [alert_summary(e, store.trace(e["alert_id"])) for e in rows]}


@app.get("/api/alert/{alert_id}")
def get_alert(alert_id: str):
    """One alert with the notice text it was read from, so the page can show where every batch number came from."""
    entry = _entry(alert_id)
    trace = store.trace(alert_id)
    return {"alert": alert_summary(entry, trace), "text": entry.get("text") or "", "trace": decorate_trace(trace, entry.get("source", ""))}


@app.get("/api/search")
def get_search(q: str = Query("", max_length=80)):
    return {"query": q, "results": search(entries(), q)}


@app.get("/api/status")
def get_status():
    """The smart update: when each regulator was last checked, what was new, how fast it was picked up."""
    return status(store.runs(500), store.state(), entries()) | {"alerts": len(entries()), "alerts_as_of": alerts_as_of()}


@app.get("/api/insights")
def get_insights():
    """Over the 90 alerts of the evaluation, all of which were traced, so the denominator stays fixed."""
    seed = {json.loads(l)["alert_id"] for l in SEED.read_text(encoding="utf-8").splitlines() if l.strip()}
    return insights(entries(), store.trace, only=seed)


@app.get("/feed.xml")
def get_feed(request: Request, source: str = Query("", max_length=60)):
    """Atom feed of new alerts, optionally for one regulator or country (?source=SFDA, ?source=Saudi)."""
    site = os.environ.get("SITE_URL") or str(request.base_url).rstrip("/")
    want = source.lower()
    xml = atom_feed(entries(), site, (lambda e: want in str(e.get("source", "")).lower()) if want else None)
    return Response(xml, media_type="application/atom+xml")


# ---------------------------------------------------------------- cross-border trace

def _entry(alert_id: str) -> dict:
    entry = store.entry(alert_id)
    if entry is None:
        raise HTTPException(404, "No alert with that id in the registry.")
    return entry


@app.get("/api/trace/{alert_id}")
def get_trace(alert_id: str):
    """A trace already run for this alert, or null. Never starts a new run."""
    _entry(alert_id)
    return {"trace": decorate_trace(store.trace(alert_id), _entry(alert_id).get("source", ""))}


def _sse(event: str, data: dict) -> str:
    return f"event: {event}\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"


@app.post("/api/trace/{alert_id}")
def post_trace(alert_id: str, request: Request):
    """Run the cross-border agent (Nemotron + Tavily) and stream its steps as server-sent events.
    A saved trace is returned at once and free; only a new run counts against the daily limit."""
    entry = _entry(alert_id)
    saved = store.trace(alert_id)
    if saved:
        return StreamingResponse(iter([_sse("done", {"trace": decorate_trace(saved, entry.get("source", ""))})]), media_type="text/event-stream")
    visitor = request.client.host if request.client else "unknown"
    refusal = limiter.take(visitor)
    if refusal:
        raise HTTPException(429, refusal)
    events: queue.Queue = queue.Queue()

    def work() -> None:
        try:
            with _trace_lock:
                result = store.trace(alert_id) or trace_cached(entry, on_step=lambda ev: events.put(("step", ev)))
            store.forget_trace(alert_id)
            events.put(("done", {"trace": decorate_trace(result, entry.get("source", ""))}))
        except SystemExit as e:  # missing API keys: nothing was spent
            limiter.give_back(visitor)
            events.put(("error", {"detail": str(e)}))
        except Exception as e:
            from batchwatch.llm import out_of_credit
            if out_of_credit(e):
                limiter.give_back(visitor)
                events.put(("error", {"detail": "The model credit for this demo is used up, so no new trace can run "
                                                "right now. Saved traces still work."}))
            else:
                events.put(("error", {"detail": f"The trace stopped: {type(e).__name__}. Try again later."}))

    threading.Thread(target=work, daemon=True).start()

    def stream():
        yield _sse("start", {"model": "Nemotron 3 Super", "tools": ["Tavily search", "Tavily extract"],
                             "limits": limiter.status()})
        while True:
            try:
                kind, data = events.get(timeout=15)
            except queue.Empty:
                yield ": still working\n\n"  # keeps proxies from closing the connection
                continue
            yield _sse(kind, data)
            if kind in ("done", "error"):
                return

    return StreamingResponse(stream(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.get("/api/limits")
def get_limits():
    return limiter.status()


app.mount("/", StaticFiles(directory=ROOT / "web", html=True), name="web")
